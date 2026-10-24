"""Điền báo cáo tuần 1 trên đúng mẫu Word của môn học."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "sources" / "mau_bao_cao.docx"
ASSIGNMENTS = ROOT / "sources" / "91_phieu.docx"
OUTPUT = ROOT / "nop_tuan_01" / "N24DCDT046_DT041_LeMinhLong_BaoCao.docx"


def set_paragraph(doc, index, value):
    doc.paragraphs[index].text = value


def set_row(table, index, values):
    for cell, value in zip(table.rows[index].cells, values):
        cell.text = value


def insert_original_assignment(report, source):
    """Sao chép nguyên phiếu DT041, gồm cả các bảng ví dụ."""
    body = source.element.body
    children = list(body.iterchildren())
    start = end = None
    for i, child in enumerate(children):
        text = "".join(node.text or "" for node in child.xpath(".//w:t"))
        if "DT041   INT13185" in text and start is None:
            # Lùi về tên cơ sở và tiêu đề phiếu.
            start = i - 3
        if start is not None and "DT042   INT13185" in text:
            end = i - 3
            break
    if start is None or end is None:
        raise RuntimeError("Không xác định được ranh giới phiếu DT041")
    anchor = report.paragraphs[31]._element
    for child in children[start:end]:
        anchor.addprevious(deepcopy(child))


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    report = Document(TEMPLATE)
    source = Document(ASSIGNMENTS)
    tables = list(report.tables)

    # Hai bìa của mẫu.
    for index in (3, 16):
        set_paragraph(report, index, "KHOA ĐIỆN TỬ 2")
    for index in (7, 20):
        set_paragraph(report, index, "Ngành Kỹ thuật điện, điện tử")
    for table in tables[:2]:
        set_row(table, 1, ["Sinh viên thực hiện", "Lê Minh Long"])
        set_row(table, 2, ["Mã số sinh viên", "N24DCDT046"])
        set_row(table, 3, ["Lớp", "D24CQDT01-N"])
    tables[1].cell(5, 1).text = "Bản tiến độ tuần 1/4, ngày 24/10/2026"
    for index in (9, 22):
        set_paragraph(report, index, "GIẢI SUDOKU THU NHỎ 4×4 BẰNG QUAY LUI")
    for index in (10, 23):
        set_paragraph(report, index, "Mã đề DT041")
    for index in (12, 25):
        set_paragraph(report, index, "THÀNH PHỐ HỒ CHÍ MINH, THÁNG 10 NĂM 2026")

    set_row(tables[2], 1, ["Họ tên, MSSV, lớp", "Lê Minh Long; N24DCDT046; D24CQDT01-N"])
    set_row(tables[2], 2, ["Mã đề và tên đề", "DT041 — Giải Sudoku thu nhỏ 4×4 bằng quay lui"])
    set_row(tables[2], 3, ["Căn cứ giáo trình", "Chương 3, mục 3.6.2–3.6.3; Bài 3.4; trang in 80–81, 84; trang PDF 92–93, 96"])
    set_row(tables[2], 4, ["Ngày giao, ngày nộp", "Theo thông báo của giảng viên trên Classroom; phiếu không điền ngày cụ thể."])

    set_paragraph(report, 63, "LỜI CAM ĐOAN")
    set_paragraph(report, 64, "Báo cáo tiến độ tuần 1 ghi đúng phần mã và kết quả đã thực hiện. Các phần solver, giao diện web và minh chứng cuối kỳ sẽ được bổ sung trong các tuần tiếp theo.")
    set_paragraph(report, 65, "Sinh viên thực hiện: Lê Minh Long — N24DCDT046")
    set_paragraph(report, 69, "Tuần 1 đã đặc tả Sudoku 4×4, xây dựng bộ kiểm tra đầu vào và kiểm tra một nước đặt số bằng Python. Hai ví dụ trên phiếu đã được nạp và kiểm tra tính hợp lệ. Bộ giải quay lui, giao diện web và minh chứng cuối kỳ đang trong kế hoạch các tuần tiếp theo.")

    content = {
        151: "Đề DT041 yêu cầu điền một bảng Sudoku 4×4 với các khối 2×2 bằng quay lui. Bản thu nhỏ giúp quan sát rõ thao tác thử số, kiểm ràng buộc và hoàn tác. Tuần 1 tập trung xác định đúng miền dữ liệu và ngăn đề sai đi vào bước tìm kiếm.",
        154: "Sản phẩm cuối gồm chương trình Python chạy terminal, bộ kiểm thử, báo cáo và kho GitHub; giao diện web được triển khai thêm theo kế hoạch cá nhân. Ở tuần 1, terminal đã in bảng, kiểm cấu trúc, miền giá trị, xung đột và kiểm một nước đặt. Chương trình chưa tạo nghiệm.",
        158: "Đầu vào là bốn hàng, mỗi hàng bốn số nguyên từ 0 đến 4; 0 là ô trống. Mỗi số cho trước khác 0 không được lặp trong hàng, cột hoặc khối 2×2. Terminal cho phép nạp hai ví dụ hoặc nhập 16 số từ bàn phím. Giao diện hiển thị chỉ số hàng/cột 1–4, lõi Python dùng 0–3.",
        161: "Bản cuối phải trả một bảng nghiệm hoàn chỉnh, thông báo không có nghiệm, hoặc báo đề không hợp lệ trước khi tìm kiếm. Nghiệm phải giữ các số cho trước. Tuần 1 mới có thông báo đề hợp lệ/không hợp lệ và kết quả kiểm một nước đặt; các trạng thái sau tìm kiếm chưa được cài.",
        167: "Quy trình gồm đọc phiếu và giáo trình, giải tay hai ví dụ, đặc tả dữ liệu, thiết kế hàm, cài đặt, kiểm thử và lưu minh chứng. Chương 1 nêu bài toán; chương 2 nêu cơ sở; chương 3 trình bày thiết kế và mã; chương 4 dành cho kết quả; chương 5 đối chiếu mức hoàn thành.",
        173: "Sudoku 4×4 gồm 16 ô; mỗi hàng, cột và khối 2×2 chứa các số 1–4 không lặp. Số 0 biểu thị ô chưa gán. Giáo trình INT1306 trình bày Sudoku như bài toán thỏa ràng buộc và quay lui theo thứ tự quét tại mục 3.6.2–3.6.3, trang in 80–81, 84.",
        177: "Bảng được lưu bằng bốn dãy con, mỗi dãy có bốn số nguyên. `validate_board` tạo bản tuple bất biến sau khi kiểm đề; `can_place` đọc bảng, không sửa dữ liệu. Bản solver tuần sau sẽ dùng bản bảng có thể thay đổi để gán và hoàn tác.",
        180: "Miền ô là 0–4. Số khác 0 ở đề gốc là số cho trước và phải được giữ nguyên. Một nước đặt chỉ xét ô đang bằng 0; giá trị định đặt phải thuộc 1–4 và không trùng trong ba vùng liên quan.",
        189: "Quay lui xây lời giải từng ô. Tại ô trống đầu tiên theo thứ tự hàng/cột, thử 1–4, kiểm hàng/cột/khối trước khi đặt. Nếu nhánh con thất bại, trả ô về 0 và thử số kế tiếp. Hai hàm kiểm đã được cài ở tuần 1; thao tác tìm kiếm và hoàn tác sẽ được cài ở tuần 2.",
        192: "Bắt đầu từ bảng đã qua kiểm tra đầu vào. Trường hợp cơ sở là không còn ô trống; khi ấy bảng hoàn chỉnh cần qua bộ kiểm nghiệm. Mỗi lời gọi tiến tới trường hợp cơ sở bằng cách giảm một ô trống, hoặc kết thúc sau tối đa bốn thử tại ô hiện tại.",
        195: "Gọi z là số ô trống. Cây thử thô của thuật toán cuối có tối đa 4^z nhánh lá; kiểm một số quét số ô cố định nên chi phí mỗi lần kiểm là hằng số với bảng 4×4. Bộ nhớ ngăn xếp dự kiến O(z). Tuần 1 chỉ đo kết quả kiểm đề, chưa có số liệu về tìm kiếm.",
        202: "`validate_board(board)` nhận bảng bất kỳ, trả tuple 4×4 nếu hợp lệ hoặc phát sinh `BoardError`. `can_place(board,row,col,value)` nhận bảng hợp lệ, chỉ số 0–3 và số 1–4, trả boolean; hàm không thay đổi bảng. Terminal chuyển tọa độ nhập 1–4 sang 0–3.",
        205: "Bộ kiểm từ chối thiếu/thừa hàng, hàng sai số ô, giá trị không phải số nguyên 0–4, trùng số cho trước trong hàng, cột và khối 2×2. Số 0 không được tính là trùng. Đề hợp lệ nhưng vô nghiệm chỉ có thể xác định sau khi bổ sung solver.",
        208: "Tuần 1 đã cài `validate_board` và `can_place` trong `src/core/validation.py`, `read_board`/`main` trong `src/terminal.py`. Lõi sẽ dùng chung cho terminal và web. Solver `solve_sudoku` cùng hàm kiểm nghiệm cuối mới ở mức thiết kế.",
        219: "GIẢI(b): tìm ô 0 đầu theo thứ tự hàng; nếu không có, kiểm nghiệm rồi trả thành công. Với v từ 1 đến 4: nếu đặt được thì gán v, gọi GIẢI(b); nếu nhánh thành công trả thành công; nếu thất bại đặt lại 0. Thử hết thì trả thất bại. Đây là thiết kế của tuần 2, chưa phải mã đã chạy.",
        254: "Môi trường tuần 1: Python 3, thư viện chuẩn `unittest`; chạy `python -m unittest discover -s tests -v`. Phiên bản Python cụ thể và SHA được ghi trong tệp ThongTinNop khi đóng gói. Không có kết quả solver ở tuần này.",
        257: "Expected của validator được suy từ quy tắc Sudoku 4×4: kiểm trực tiếp các số cho trước trong hàng, cột, khối. Hai ví dụ trên phiếu được dùng làm dữ liệu đầu vào hợp lệ, chưa dùng để khẳng định thuật toán giải đã đạt nghiệm.",
        260: "Bảng kế hoạch 16 ca ở `docs/ke_hoach_kiem_thu.md`. Tuần 1 đã chạy các nhóm kiểm cấu trúc, miền giá trị, xung đột và `can_place`. Hai ví dụ chỉ được xác nhận là đầu vào hợp lệ; kết quả giải và tổng số ca độc lập cuối kỳ sẽ được điền sau khi chạy solver.",
    }
    for index, value in content.items():
        set_paragraph(report, index, value)
    for index in range(151, 251):
        paragraph = report.paragraphs[index]
        if paragraph.style.name == "Nội dung" and index not in content:
            paragraph.text = ""

    # Mẫu chứa lời nhắc điền số liệu và ảnh của kỳ nộp cuối. Tuần 1 để các
    # ô chưa thực hiện trống theo hướng dẫn nộp tiến độ, tránh biến lời nhắc
    # thành kết quả được khẳng định trong báo cáo.
    for index in range(261, 400):
        paragraph = report.paragraphs[index]
        if paragraph.style.name not in ("Heading 1", "Heading 2", "Heading 3"):
            paragraph.text = ""
    set_paragraph(report, 364, "Giáo trình Cấu trúc dữ liệu và Giải thuật INT1306, chương 3, mục 3.6.2–3.6.3 và Bài 3.4.")
    set_paragraph(report, 370, "Mã nguồn, lệnh chạy và hướng dẫn kiểm thử nằm trong README.md của tệp MaNguon.zip.")
    set_paragraph(report, 380, "Tuần 1: đã đối chiếu thông tin cá nhân, phiếu DT041, chạy 7 bài kiểm thử Python; các mục cuối kỳ còn lại sẽ được kiểm tra ở tuần 4.")

    set_row(tables[6], 1, ["Bảng b", "4 hàng × 4 ô", "Số nguyên 0–4", "Bàn phím hoặc khai báo"])
    set_row(tables[6], 2, ["Ô trống", "Một ô b[r][c]", "Giá trị 0", "Từ bảng b"])
    set_row(tables[6], 3, ["Số định đặt", "Một số nguyên", "1–4, ô đích trống", "Tham số --check"])
    set_row(tables[7], 1, ["Tình trạng đề", "Thông báo terminal", "Hợp lệ hoặc sai quy tắc", "Đề sai: từ chối trước tìm kiếm"])
    set_row(tables[7], 2, ["Nghiệm cuối", "Bảng 4×4", "Một nghiệm giữ ô cho trước", "Chưa cài trong tuần 1"])
    set_row(tables[8], 1, ["Kiểm dữ liệu đầu vào", "Đã cài và kiểm thử", "Chương 3, 4"])
    set_row(tables[8], 2, ["Giải bằng quay lui", "Thiết kế xong; cài ở tuần 2", "Chương 3"])
    set_row(tables[8], 3, ["Kiểm thử và minh chứng", "Có kế hoạch 16 ca; hoàn thiện ở tuần 2–4", "Chương 4"])
    set_row(tables[9], 1, ["b[r][c]", "Giá trị ô", "0–4; ô cho trước giữ nguyên"])
    set_row(tables[9], 2, ["z", "Số ô trống", "0–16; giảm khi gọi sâu"])
    set_row(tables[10], 1, ["validate_board", "Bảng bất kỳ", "Tuple bảng hoặc BoardError", "Đã cài; kiểm đề"])
    set_row(tables[10], 2, ["can_place", "Bảng, r, c, v", "Boolean", "Đã cài; không sửa bảng"])
    set_row(tables[10], 3, ["solve_sudoku", "Bảng hợp lệ", "Nghiệm hoặc vô nghiệm", "Thiết kế; tuần 2"])
    set_row(tables[11], 1, ["validate_board(board)", "Bảng có thể sai", "Trả bản bất biến nếu hợp lệ", "BoardError có giải thích"])
    set_row(tables[11], 2, ["can_place(board,r,c,v)", "Đề hợp lệ; r,c 0–3; v 1–4", "Boolean; bảng không đổi", "ValueError nếu tham số sai"])
    set_row(tables[11], 3, ["solve_sudoku(board)", "Đề hợp lệ", "Một nghiệm hoặc không có nghiệm", "Chưa cài trong tuần 1"])

    # Các bảng kết quả chưa làm để trống theo yêu cầu nộp tiến độ của giảng viên.
    for table in tables[12:31]:
        for row in table.rows[1:]:
            for cell in row.cells:
                if "[" in cell.text and "]" in cell.text:
                    cell.text = ""

    # Chèn phiếu sau các sửa theo vị trí của mẫu để chỉ số đoạn không bị dịch.
    insert_original_assignment(report, source)
    report.core_properties.title = "Báo cáo tiến độ tuần 1 DT041 Sudoku 4×4"
    report.core_properties.author = "Lê Minh Long"
    report.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

